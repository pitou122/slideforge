"""Stage 5 — PPTX Generation.

Rebuild the slide deck from the extracted JSON using python-pptx.
See pipeline_spec.md for the authoritative spec.
"""
from __future__ import annotations

import math
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_ANCHOR
from pptx.util import Inches, Pt

# Slide is 10" x 5.625" (16:9).
SLIDE_W = 10.0
SLIDE_H = 5.625

# Column X positions and widths per layout (inches).
LAYOUTS = {
    "single": {"x": [0.5], "w": [9.0]},
    "two_column": {"x": [0.5, 5.25], "w": [4.5, 4.5]},
    "three_column": {"x": [0.5, 3.67, 6.84], "w": [2.9, 2.9, 2.9]},
}

FONT_PT = {"title": 28, "subtitle": 20, "body": 12, "small": 10, "tiny": 8}

TITLE_BOX = {"x": 0.3, "y": 0.15, "w": 9.4, "h": 0.9}
CONTENT_TOP = 1.2     # y where column content begins
ELEMENT_GAP = 0.1     # vertical gap between stacked elements
BOTTOM_LIMIT = 5.3    # skip an element whose bottom edge would exceed this


def _rgb(hex_str: str | None, default: str = "000000") -> RGBColor:
    """Parse a ``#RRGGBB`` string into an RGBColor, falling back to ``default``."""
    s = (hex_str or "").lstrip("#").strip()
    if len(s) == 6:
        try:
            return RGBColor.from_string(s.upper())
        except ValueError:
            pass
    return RGBColor.from_string(default)


def _resolve_layout(layout: str, columns: list) -> dict:
    """Return a column-geometry dict sized to the actual column data.

    The model's ``layout`` string can disagree with the number of columns it
    returns (e.g. ``layout="single"`` with two columns). Trusting the string
    would clamp every column to index 0 and render them on top of each other,
    so geometry is derived from the real column count / indices instead.
    """
    if columns:
        n = max(
            len(columns),
            max((c.get("index", 0) for c in columns), default=0) + 1,
        )
    else:
        n = {"two_column": 2, "three_column": 3}.get(layout, 1)
    if n >= 3:
        return LAYOUTS["three_column"]
    if n == 2:
        return LAYOUTS["two_column"]
    return LAYOUTS["single"]


def _estimate_text_height(text: str, font_pt: int, width_in: float) -> float:
    """Rough height (inches) for word-wrapped ``text`` in a box of ``width_in``."""
    text = text or ""
    # Approximate average glyph width as 0.55 * font size.
    char_w_in = (font_pt * 0.55) / 72.0
    chars_per_line = max(1, int(width_in / char_w_in))
    explicit_lines = text.count("\n") + 1
    wrapped_lines = sum(
        max(1, math.ceil(len(seg) / chars_per_line)) for seg in text.split("\n")
    )
    lines = max(explicit_lines, wrapped_lines)
    line_h_in = (font_pt * 1.25) / 72.0
    return lines * line_h_in + 0.12  # small padding for the textbox insets


def _add_text(slide, el: dict, x: float, y: float, w: float, *, default_color: str):
    """Add a textbox for a text element; return its height in inches."""
    size_key = el.get("size", "body")
    font_pt = FONT_PT.get(size_key, FONT_PT["body"])
    text = el.get("text", "") or ""

    height = _estimate_text_height(text, font_pt, w)
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(height))
    tf = box.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.TOP

    para = tf.paragraphs[0]
    run = para.add_run()
    run.text = text
    run.font.size = Pt(font_pt)
    run.font.bold = bool(el.get("bold", False))
    run.font.color.rgb = _rgb(el.get("color"), default_color)

    bg = el.get("background")
    if bg:
        box.fill.solid()
        box.fill.fore_color.rgb = _rgb(bg, "FFFFFF")
    return height


def _add_image(slide, el: dict, out_dir: Path, x: float, y: float,
               w: float | None = None, h: float | None = None):
    """Add an image_crop picture; pass w or h (or both) to control size."""
    crop_path = el.get("crop_path")
    if not crop_path:
        return None
    full = out_dir / crop_path
    if not full.exists():
        return None
    kwargs: dict = {}
    if w is not None:
        kwargs["width"] = Inches(w)
    if h is not None:
        kwargs["height"] = Inches(h)
    slide.shapes.add_picture(str(full), Inches(x), Inches(y), **kwargs)
    return h


def build_pptx(extracted: list[dict], out_dir: str | Path) -> Path:
    """Build ``output.pptx`` from the extracted slide list.

    ``extracted`` is the Stage 3/4 output (with ``crop_path`` fields).
    Returns the path to the written ``.pptx`` file.
    """
    out_dir = Path(out_dir)

    prs = Presentation()
    prs.slide_width = Inches(SLIDE_W)
    prs.slide_height = Inches(SLIDE_H)
    blank_layout = prs.slide_layouts[6]  # fully blank layout

    for slide_data in extracted:
        slide = prs.slides.add_slide(blank_layout)

        # Background fill.
        bg_hex = slide_data.get("background_color")
        if bg_hex:
            slide.background.fill.solid()
            slide.background.fill.fore_color.rgb = _rgb(bg_hex, "FFFFFF")

        # Title.
        title = slide_data.get("title")
        if title and title.get("text"):
            box = slide.shapes.add_textbox(
                Inches(TITLE_BOX["x"]), Inches(TITLE_BOX["y"]),
                Inches(TITLE_BOX["w"]), Inches(TITLE_BOX["h"]),
            )
            tf = box.text_frame
            tf.word_wrap = True
            run = tf.paragraphs[0].add_run()
            run.text = title["text"]
            run.font.size = Pt(FONT_PT.get(title.get("size", "title"), 28))
            run.font.bold = bool(title.get("bold", True))
            run.font.color.rgb = _rgb(title.get("color"), "E07B00")

        columns = slide_data.get("columns", []) or []
        geom = _resolve_layout(slide_data.get("layout", "single"), columns)
        n_cols = len(geom["x"])

        for column in columns:
            idx = column.get("index", 0)
            if idx >= n_cols:
                idx = n_cols - 1  # clamp extra columns into the last slot
            col_x = geom["x"][idx]
            col_w = geom["w"][idx]

            cursor_y = CONTENT_TOP
            for el in column.get("elements", []):
                el_type = el.get("type")

                if el_type == "image_crop":
                    crop_path = el.get("crop_path")
                    if not crop_path or not (out_dir / crop_path).exists():
                        continue
                    from PIL import Image as _Image
                    with _Image.open(out_dir / crop_path) as im:
                        iw, ih = im.size
                    if not iw:
                        continue
                    max_h = BOTTOM_LIMIT - cursor_y - 0.05
                    if max_h < 0.2:
                        continue
                    h_at_full_w = col_w * ih / iw
                    if h_at_full_w <= max_h:
                        # Fits at full column width — scale by width only.
                        _add_image(slide, el, out_dir, col_x, cursor_y, w=col_w)
                        height = h_at_full_w
                    else:
                        # Too tall — constrain by available height, let width shrink.
                        _add_image(slide, el, out_dir, col_x, cursor_y, h=max_h)
                        height = max_h
                else:
                    size_key = el.get("size", "body")
                    font_pt = FONT_PT.get(size_key, FONT_PT["body"])
                    height = _estimate_text_height(
                        el.get("text", ""), font_pt, col_w
                    )
                    if cursor_y + height > BOTTOM_LIMIT:
                        continue
                    _add_text(
                        slide, el, col_x, cursor_y, col_w,
                        default_color="000000",
                    )

                cursor_y += height + ELEMENT_GAP

    out_path = out_dir / "output.pptx"
    prs.save(str(out_path))
    return out_path
